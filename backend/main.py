"""
DocxStudio Backend - FastAPI Application
Web-based DOCX editor with style preservation and AI assistance.
"""
import uuid
import io
from datetime import datetime, timedelta
from typing import Optional
from contextlib import asynccontextmanager

import docx
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse

from models.schemas import (
    UploadResponse, ReplaceRequest, ReplaceResponse,
    StyleRequest, StyleResponse, AiSuggestRequest, AiSuggestResponse
)
from services.parser import parse_document
from services.editor import replace_in_document
from services.styler import apply_style_changes
from services.ai_assistant import get_ai_suggestion


# In-memory session storage with TTL
sessions: dict[str, dict] = {}
SESSION_TTL_MINUTES = 30


def cleanup_old_sessions():
    """Remove sessions older than TTL."""
    now = datetime.utcnow()
    expired = [
        sid for sid, data in sessions.items()
        if now - data["created_at"] > timedelta(minutes=SESSION_TTL_MINUTES)
    ]
    for sid in expired:
        del sessions[sid]


@asynccontextmanager
async def lifespan(app: FastAPI):
    """Application lifespan handler."""
    yield
    sessions.clear()


app = FastAPI(
    title="DocxStudio API",
    description="Web-based DOCX editor with style preservation and AI assistance",
    version="1.0.0",
    lifespan=lifespan
)

# CORS configuration
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/api/health")
async def health_check():
    """Health check endpoint."""
    return {"status": "ok", "sessions": len(sessions)}


@app.post("/api/upload", response_model=UploadResponse)
async def upload_document(file: UploadFile = File(...)):
    """
    Upload a DOCX file for editing.
    Returns parsed document structure and HTML preview.
    """
    if not file.filename.endswith('.docx'):
        raise HTTPException(400, "Only .docx files are supported")
    
    cleanup_old_sessions()
    
    content = await file.read()
    
    try:
        parsed = parse_document(content)
    except Exception as e:
        raise HTTPException(400, f"Failed to parse document: {str(e)}")
    
    session_id = str(uuid.uuid4())
    sessions[session_id] = {
        "filename": file.filename,
        "original_content": content,
        "current_content": content,
        "created_at": datetime.utcnow()
    }
    
    return UploadResponse(
        session_id=session_id,
        filename=file.filename,
        paragraphs=parsed["paragraphs"],
        stats=parsed["stats"],
        html_preview=parsed["html_preview"]
    )


@app.post("/api/replace", response_model=ReplaceResponse)
async def replace_text(request: ReplaceRequest):
    """
    Find and replace text in the document.
    Preserves original formatting.
    """
    if request.session_id not in sessions:
        raise HTTPException(404, "Session not found or expired")
    
    session = sessions[request.session_id]
    
    doc = docx.Document(io.BytesIO(session["current_content"]))
    
    replacements = [{"find": r.find, "replace": r.replace} for r in request.replacements]
    result = replace_in_document(doc, replacements, request.case_sensitive)
    
    # Save modified document back to session
    output = io.BytesIO()
    doc.save(output)
    session["current_content"] = output.getvalue()
    session["created_at"] = datetime.utcnow()
    
    return ReplaceResponse(
        total_replacements=result["total_replacements"],
        affected_paragraphs=result["affected_paragraphs"],
        html_preview=result["html_preview"]
    )


@app.post("/api/style", response_model=StyleResponse)
async def change_styles(request: StyleRequest):
    """
    Modify styles in the document based on matching criteria.
    Example: Change all bold red text to purple.
    """
    if request.session_id not in sessions:
        raise HTTPException(404, "Session not found or expired")
    
    session = sessions[request.session_id]
    
    doc = docx.Document(io.BytesIO(session["current_content"]))
    
    changes = [
        {
            "match": c.match.model_dump(exclude_none=True),
            "apply": c.apply.model_dump(exclude_none=True)
        }
        for c in request.changes
    ]
    result = apply_style_changes(doc, changes)
    
    # Save modified document back to session
    output = io.BytesIO()
    doc.save(output)
    session["current_content"] = output.getvalue()
    session["created_at"] = datetime.utcnow()
    
    return StyleResponse(
        total_changes=result["total_changes"],
        affected_paragraphs=result["affected_paragraphs"],
        html_preview=result["html_preview"]
    )


@app.post("/api/ai-suggest", response_model=AiSuggestResponse)
async def ai_suggest(request: AiSuggestRequest):
    """
    Get AI-powered suggestions for document editing.
    Describe changes in natural language.
    """
    if request.session_id not in sessions:
        raise HTTPException(404, "Session not found or expired")
    
    session = sessions[request.session_id]
    
    # Get document text for context
    doc = docx.Document(io.BytesIO(session["current_content"]))
    doc_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()][:20])
    
    result = await get_ai_suggestion(request.message, doc_text)
    
    return AiSuggestResponse(
        suggestion=result["suggestion"],
        explanation=result["explanation"]
    )


@app.get("/api/download/{session_id}")
async def download_document(session_id: str):
    """
    Download the modified DOCX document.
    """
    if session_id not in sessions:
        raise HTTPException(404, "Session not found or expired")
    
    session = sessions[session_id]
    content = session["current_content"]
    filename = session["filename"].replace(".docx", "_edited.docx")
    
    return StreamingResponse(
        io.BytesIO(content),
        media_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@app.delete("/api/session/{session_id}")
async def delete_session(session_id: str):
    """Delete a session and free memory."""
    if session_id in sessions:
        del sessions[session_id]
    return {"status": "deleted"}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
